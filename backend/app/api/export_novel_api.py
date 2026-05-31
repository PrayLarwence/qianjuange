"""风格档案 / 小说化导出（export-novel）API。

端点：
- GET    /style_profiles
- GET    /style_profiles/{id}
- POST   /style_profiles              (从 builtin 派生 custom)
- PATCH  /style_profiles/{id}         (更新 custom 的结构化字段)
- POST   /style_profiles/{id}/reset   (恢复为 builtin 默认值)
- DELETE /style_profiles/{id}         (删除 custom)
- POST   /worlds/{id}/export
"""
from __future__ import annotations
import io
import json
import logging
import uuid
from datetime import datetime
from typing import Any, Optional
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ..models import (
    get_db, World, Branch, Entity, Event, NarrativeLog, ChapterMarker, StyleProfile,
)

log = logging.getLogger(__name__)
router = APIRouter()


# ============== style profiles (A4) ==============

@router.get("/style_profiles")
def list_style_profiles(db: Session = Depends(get_db)):
    """列所有风格档案。前端在'世界设置'弹窗里展示给用户选择。

    返回 builtin（系统内置）+ custom（用户自建）两组。spec_text 不返回——
    前端只需要展示用元数据，需要看 spec 再单独 GET /style_profiles/{id}。
    """
    rows = db.query(StyleProfile).order_by(StyleProfile.kind.asc(), StyleProfile.name.asc()).all()
    return {
        "profiles": [
            {
                "id": r.id,
                "name": r.name,
                "description": r.description or "",
                "kind": r.kind or "builtin",
                "category": r.category or "",
                "frozen": bool(r.frozen),
            }
            for r in rows
        ]
    }


@router.get("/style_profiles/{style_id}")
def get_style_profile(style_id: str, db: Session = Depends(get_db)):
    r = db.query(StyleProfile).filter_by(id=style_id).first()
    if not r:
        raise HTTPException(404, "style_profile not found")
    return {
        "id": r.id,
        "name": r.name,
        "description": r.description or "",
        "kind": r.kind or "builtin",
        "category": r.category or "",
        "spec_text": r.spec_text or "",
        "sample_paragraphs": r.sample_paragraphs or [],
        "negative_rules": r.negative_rules or [],
        "builtin_source_id": r.builtin_source_id or "",
        "frozen": bool(r.frozen),
    }


class NegativeRuleItem(BaseModel):
    label: str
    description: str
    enabled: bool = True


class CreateStyleProfileRequest(BaseModel):
    source_id: str
    name: str
    world_id: Optional[str] = None


@router.post("/style_profiles")
def create_style_profile(payload: CreateStyleProfileRequest, db: Session = Depends(get_db)):
    """从一个已有 profile（通常是 builtin）派生一个 custom 副本。"""
    source = db.query(StyleProfile).filter_by(id=payload.source_id).first()
    if not source:
        raise HTTPException(404, f"source profile not found: {payload.source_id}")
    new_id = f"style_custom_{uuid.uuid4().hex[:8]}"
    profile = StyleProfile(
        id=new_id,
        name=payload.name.strip() or f"{source.name}（自定义）",
        kind="custom",
        category=source.category or "",
        description=source.description or "",
        world_id=payload.world_id,
        spec_text=source.spec_text or "",
        sample_paragraphs=source.sample_paragraphs or [],
        negative_rules=source.negative_rules or [],
        builtin_source_id=source.id if source.kind == "builtin" else (source.builtin_source_id or source.id),
        frozen=0,
        created_at=datetime.utcnow(),
        updated_at=datetime.utcnow(),
    )
    db.add(profile)
    db.commit()
    return {
        "id": profile.id,
        "name": profile.name,
        "kind": "custom",
        "builtin_source_id": profile.builtin_source_id,
    }


class UpdateStyleProfileRequest(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    spec_text: Optional[str] = None
    sample_paragraphs: Optional[list[dict]] = None
    negative_rules: Optional[list[NegativeRuleItem]] = None


@router.patch("/style_profiles/{style_id}")
def update_style_profile(style_id: str, payload: UpdateStyleProfileRequest, db: Session = Depends(get_db)):
    """更新 custom profile 的字段。builtin 不可修改。"""
    r = db.query(StyleProfile).filter_by(id=style_id).first()
    if not r:
        raise HTTPException(404, "style_profile not found")
    if r.kind == "builtin":
        raise HTTPException(403, "内置风格不可修改，请先派生为自定义副本")
    if r.frozen:
        raise HTTPException(403, "该风格已冻结")
    if payload.name is not None:
        r.name = payload.name.strip() or r.name
    if payload.description is not None:
        r.description = payload.description
    if payload.spec_text is not None:
        r.spec_text = payload.spec_text
    if payload.sample_paragraphs is not None:
        r.sample_paragraphs = payload.sample_paragraphs
    if payload.negative_rules is not None:
        r.negative_rules = [item.model_dump() for item in payload.negative_rules]
    r.updated_at = datetime.utcnow()
    db.commit()
    return {"ok": True, "id": r.id}


@router.post("/style_profiles/{style_id}/reset")
def reset_style_profile(style_id: str, db: Session = Depends(get_db)):
    """将 custom profile 恢复为其 builtin 源的默认值。"""
    r = db.query(StyleProfile).filter_by(id=style_id).first()
    if not r:
        raise HTTPException(404, "style_profile not found")
    if r.kind == "builtin":
        raise HTTPException(400, "内置风格无需恢复")
    source_id = r.builtin_source_id
    if not source_id:
        raise HTTPException(400, "该自定义风格没有关联的内置源，无法恢复")
    source = db.query(StyleProfile).filter_by(id=source_id).first()
    if not source:
        raise HTTPException(404, f"内置源 {source_id} 不存在")
    r.spec_text = source.spec_text
    r.sample_paragraphs = source.sample_paragraphs
    r.negative_rules = source.negative_rules
    r.description = source.description
    r.updated_at = datetime.utcnow()
    db.commit()
    return {"ok": True, "id": r.id, "reset_from": source_id}


@router.delete("/style_profiles/{style_id}")
def delete_style_profile(style_id: str, db: Session = Depends(get_db)):
    """删除 custom profile。builtin 不可删。"""
    r = db.query(StyleProfile).filter_by(id=style_id).first()
    if not r:
        raise HTTPException(404, "style_profile not found")
    if r.kind == "builtin":
        raise HTTPException(403, "内置风格不可删除")
    db.query(World).filter_by(style_profile_id=style_id).update(
        {"style_profile_id": None}, synchronize_session=False
    )
    db.delete(r)
    db.commit()
    return {"ok": True}



class ExportRequest(BaseModel):
    mode: str = "raw"  # raw|json|novelize|docx|epub
    tick_from: int | None = None
    tick_to: int | None = None
    chapter_size: int = 5
    use_chapter_markers: bool = True
    include_critique: bool = False
    include_events: bool = True
    include_narration: bool = True
    include_entities: bool = False
    provider: str | None = None


@router.post("/worlds/{world_id}/export")
def export_world(world_id: str, payload: ExportRequest, db: Session = Depends(get_db)):
    world = db.query(World).filter_by(id=world_id).first()
    if not world:
        raise HTTPException(404, "world not found")
    branch_id = world.active_branch_id
    branch = db.query(Branch).filter_by(id=branch_id).first()

    tick_lo = 0 if payload.tick_from is None else max(0, payload.tick_from)
    tick_hi = world.current_tick if payload.tick_to is None else payload.tick_to

    events = (
        db.query(Event).filter_by(branch_id=branch_id, deleted=0)
        .filter(Event.tick >= tick_lo, Event.tick <= tick_hi)
        .order_by(Event.tick).all()
    )
    narration = (
        db.query(NarrativeLog).filter_by(branch_id=branch_id)
        .filter(NarrativeLog.tick >= tick_lo, NarrativeLog.tick <= tick_hi)
        .order_by(NarrativeLog.tick).all()
    )
    entities = db.query(Entity).filter_by(branch_id=branch_id).all()
    by_id = {e.id: e for e in entities}

    if payload.mode == "json":
        return {
            "format": "json",
            "filename": f"{world.name}_{branch.name if branch else 'main'}_{tick_lo}-{tick_hi}.json",
            "content": json.dumps({
                "world": {"name": world.name, "description": world.description, "rules": world.rules or {}},
                "branch": branch.name if branch else "main",
                "tick_range": [tick_lo, tick_hi],
                "entities": [{"id": e.id, "type": e.type, "name": e.name, "summary": e.summary,
                              "attributes": e.attributes or {}, "state": e.state or {}, "alive": e.alive,
                              "location_id": e.location_id, "created_at_tick": e.created_at_tick} for e in entities],
                "events": [{"id": ev.id, "tick": ev.tick, "title": ev.title, "description": ev.description,
                            "participants": ev.participants or [], "consequences": ev.consequences or [],
                            "location_id": ev.location_id} for ev in events],
                "narration": [{"tick": n.tick, "role": n.role, "text": n.text} for n in narration],
            }, ensure_ascii=False, indent=2),
        }

    if payload.mode == "raw":
        lines = [f"# {world.name}", ""]
        if world.description:
            lines += [world.description, ""]
        lines += [f"*分支：{branch.name if branch else 'main'} · tick {tick_lo}–{tick_hi}*", ""]

        if payload.include_entities:
            chars = [e for e in entities if e.type == "character"]
            if chars:
                lines += ["## 主要人物", ""]
                for e in chars:
                    state = "" if e.alive else "（已逝）"
                    lines.append(f"- **{e.name}**{state} — {e.summary or ''}")
                lines.append("")

        ticks = sorted(set([ev.tick for ev in events] + [n.tick for n in narration]))
        if ticks:
            lines += ["## 故事正文", ""]
            ev_by_tick: dict[int, list[Event]] = {}
            for ev in events:
                ev_by_tick.setdefault(ev.tick, []).append(ev)
            nar_by_tick: dict[int, list[NarrativeLog]] = {}
            for n in narration:
                nar_by_tick.setdefault(n.tick, []).append(n)

            for t in ticks:
                lines.append(f"### t{t}")
                if payload.include_narration:
                    for n in nar_by_tick.get(t, []):
                        lines.append(n.text)
                        lines.append("")
                if payload.include_events:
                    for ev in ev_by_tick.get(t, []):
                        parts = "、".join(by_id[p].name for p in (ev.participants or []) if p in by_id)
                        loc = by_id.get(ev.location_id).name if ev.location_id and ev.location_id in by_id else ""
                        meta_bits = " ".join(f"[{x}]" for x in [parts, loc] if x)
                        lines.append(f"- **{ev.title}** {meta_bits}")
                        if ev.description:
                            lines.append(f"  > {ev.description}")
                lines.append("")

        return {
            "format": "markdown",
            "filename": f"{world.name}_{tick_lo}-{tick_hi}.md",
            "content": "\n".join(lines),
        }

    if payload.mode == "docx":
        return _export_docx(world, events, narration, entities, by_id,
                            tick_lo, tick_hi, branch, payload)

    if payload.mode == "epub":
        return _export_epub(world, events, narration, entities, by_id,
                            tick_lo, tick_hi, branch, payload)

    if payload.mode == "novelize":
        # Assemble from author_final logs — no LLM rewriting.
        # Quality comes from the simulation pipeline (Director → Author → Critics).
        author_finals = (
            db.query(NarrativeLog)
            .filter(
                NarrativeLog.branch_id == branch_id,
                NarrativeLog.role == "author_final",
                NarrativeLog.tick >= tick_lo,
                NarrativeLog.tick <= tick_hi,
            )
            .order_by(NarrativeLog.tick.asc(), NarrativeLog.created_at.desc())
            .all()
        )
        # Keep only the latest revision per tick
        finals_by_tick: dict[int, NarrativeLog] = {}
        for af in author_finals:
            if af.tick not in finals_by_tick:
                finals_by_tick[af.tick] = af

        if not finals_by_tick:
            raise HTTPException(400, "no author_final logs in range — run simulation with orchestrated mode first")

        markers = []
        if payload.use_chapter_markers:
            markers = (
                db.query(ChapterMarker).filter_by(branch_id=branch_id)
                .filter(ChapterMarker.tick >= tick_lo, ChapterMarker.tick <= tick_hi)
                .order_by(ChapterMarker.tick).all()
            )

        # Build chapter boundaries
        chunks: list[tuple[int, int, str]] = []  # (tick_start, tick_end, title)
        if markers:
            cursor = tick_lo
            for m in markers:
                if cursor <= m.tick:
                    chunks.append((cursor, m.tick, m.title or ""))
                cursor = m.tick + 1
            if cursor <= tick_hi:
                chunks.append((cursor, tick_hi, ""))
        else:
            chunk_size = max(1, min(payload.chapter_size, 20))
            for chunk_start in range(tick_lo, tick_hi + 1, chunk_size):
                chunk_end = min(chunk_start + chunk_size - 1, tick_hi)
                chunks.append((chunk_start, chunk_end, ""))

        chapters_md: list[str] = [f"# {world.name}"]
        if world.description:
            chapters_md += ["", world.description]
        chunk_source = "章节标记" if markers else f"每 {payload.chapter_size} ticks"
        chapters_md += [f"\n*分支 {branch.name if branch else 'main'} · 共 {len(chunks)} 章 · 切分方式：{chunk_source}*", ""]

        for idx, (cs, ce, hint_title) in enumerate(chunks, start=1):
            title = hint_title or f"第 {idx} 章"
            chapter_texts: list[str] = []

            for t in range(cs, ce + 1):
                if t in finals_by_tick:
                    text = (finals_by_tick[t].text or "").strip()
                    if text:
                        chapter_texts.append(text)

            if not chapter_texts:
                continue

            parts = [f"## {title}", ""]
            parts.append("\n\n".join(chapter_texts))

            chapters_md.append("\n".join(parts))
            chapters_md.append("")

        result = {
            "format": "markdown",
            "filename": f"{world.name}_小说_{tick_lo}-{tick_hi}.md",
            "content": "\n".join(chapters_md),
            "chapters": len(chunks),
            "chunked_by": "markers" if markers else "size",
        }
        return result

    raise HTTPException(400, f"unknown mode: {payload.mode}")


# ─── DOCX 导出 ──────────────────────────────────────────────────

def _export_docx(world, events, narration, entities, by_id, tick_lo, tick_hi, branch, payload):
    from docx import Document
    from docx.shared import Pt, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'SimSun'
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5

    # 标题页
    title = doc.add_heading(world.name or "未命名", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if world.description:
        p = doc.add_paragraph(world.description)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"分支：{branch.name if branch else 'main'} · tick {tick_lo}–{tick_hi}").italic = True
    doc.add_page_break()

    if payload.include_entities:
        chars = [e for e in entities if e.type == "character"]
        if chars:
            doc.add_heading("主要人物", level=1)
            for e in chars:
                state = "" if e.alive else "（已逝）"
                p = doc.add_paragraph()
                p.add_run(f"• {e.name}").bold = True
                p.add_run(f"{state}  {e.summary or ''}")

    doc.add_heading("正文", level=1)

    ticks = sorted(set([ev.tick for ev in events] + [n.tick for n in narration]))
    ev_by_tick = {}
    for ev in events:
        ev_by_tick.setdefault(ev.tick, []).append(ev)
    nar_by_tick = {}
    for n in narration:
        nar_by_tick.setdefault(n.tick, []).append(n)

    for t in ticks:
        if payload.include_narration:
            for n in nar_by_tick.get(t, []):
                if n.text and n.text.strip():
                    doc.add_paragraph(n.text.strip())
        if payload.include_events:
            for ev in ev_by_tick.get(t, []):
                parts = "、".join(by_id[p].name for p in (ev.participants or []) if p in by_id)
                loc = by_id.get(ev.location_id).name if ev.location_id and ev.location_id in by_id else ""
                meta = " ".join(f"[{x}]" for x in [parts, loc] if x)
                p = doc.add_paragraph()
                p.add_run(f"■ {ev.title} {meta}").bold = True
                if ev.description:
                    doc.add_paragraph(f"  {ev.description}")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{world.name}_{tick_lo}-{tick_hi}.docx"'},
    )


# ─── EPUB 导出 ──────────────────────────────────────────────────

def _export_epub(world, events, narration, entities, by_id, tick_lo, tick_hi, branch, payload):
    import zipfile
    import uuid as _uuid

    book_id = f"urn:uuid:{_uuid.uuid4()}"
    safe_name = world.name.replace(" ", "_").replace("/", "_")[:40]

    ticks = sorted(set([ev.tick for ev in events] + [n.tick for n in narration]))
    ev_by_tick = {}
    for ev in events:
        ev_by_tick.setdefault(ev.tick, []).append(ev)
    nar_by_tick = {}
    for n in narration:
        nar_by_tick.setdefault(n.tick, []).append(n)

    # 构建 XHTML body
    body_lines = [f'<h1>{world.name or "未命名"}</h1>']
    if world.description:
        body_lines.append(f'<p><em>{world.description}</em></p>')

    if payload.include_entities:
        chars = [e for e in entities if e.type == "character"]
        if chars:
            body_lines.append('<h2>主要人物</h2><ul>')
            for e in chars:
                state = "" if e.alive else "（已逝）"
                body_lines.append(f'<li><strong>{e.name}</strong>{state} — {e.summary or ""}</li>')
            body_lines.append('</ul>')

    body_lines.append('<h2>正文</h2>')
    for t in ticks:
        if payload.include_narration:
            for n in nar_by_tick.get(t, []):
                if n.text and n.text.strip():
                    body_lines.append(f'<p>{n.text.strip()}</p>')
        if payload.include_events:
            for ev in ev_by_tick.get(t, []):
                body_lines.append(f'<p><em>■ {ev.title}</em></p>')
                if ev.description:
                    body_lines.append(f'<p>{ev.description}</p>')

    body_html = "\n".join(body_lines)

    chapter_xhtml = f'''<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" xml:lang="zh-CN">
<head><title>{world.name}</title></head>
<body>{body_html}</body>
</html>'''

    container_xml = '''<?xml version="1.0" encoding="utf-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>'''

    content_opf = f'''<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book-id">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="book-id">{book_id}</dc:identifier>
    <dc:title>{world.name}</dc:title>
    <dc:language>zh-CN</dc:language>
  </metadata>
  <manifest>
    <item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine><itemref idref="chapter"/></spine>
</package>'''

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        zf.writestr('mimetype', 'application/epub+zip', compress_type=zipfile.ZIP_STORED)
        zf.writestr('META-INF/container.xml', container_xml)
        zf.writestr('OEBPS/content.opf', content_opf)
        zf.writestr('OEBPS/chapter.xhtml', chapter_xhtml)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/epub+zip",
        headers={"Content-Disposition": f'attachment; filename="{safe_name}_{tick_lo}-{tick_hi}.epub"'},
    )

